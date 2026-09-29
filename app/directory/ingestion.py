from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any


VERIFICATION_VALUES = {"verified", "yes", "y", "true", "confirmed", "right"}


def _clean(value: Any) -> str | None:
    if value is None:
        return None
    result = re.sub(r"\s+", " ", str(value)).strip()
    return result or None


def _key(prefix: str, *parts: Any) -> str:
    normalized = "|".join((_clean(part) or "").casefold() for part in parts)
    return f"{prefix}:{hashlib.sha256(normalized.encode('utf-8')).hexdigest()}"


def _verification(value: Any, source: Any = None) -> str:
    normalized = (_clean(value) or "").casefold()
    source_text = (_clean(source) or "").casefold()
    if "conflict" in normalized:
        return "conflicting"
    if "not verified" in normalized or normalized in {"no", "n", "false"}:
        return "unverified"
    if normalized in VERIFICATION_VALUES or "verified" in normalized:
        return "verified"
    if "not verified" in source_text or "desk research" in source_text:
        return "unverified"
    if "conflict" in normalized:
        return "conflicting"
    return "unknown"


def _boolean(value: Any) -> bool | None:
    normalized = (_clean(value) or "").casefold()
    if normalized in {"yes", "y", "true", "1", "halal"}:
        return True
    if normalized in {"no", "n", "false", "0", "not halal"}:
        return False
    return None


def _integer(value: Any) -> int | None:
    """Convert workbook numeric cells; placeholders like '-' become null."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else None
    normalized = (_clean(value) or "").replace(",", "")
    if not normalized or normalized in {"-", "—", "–", "n/a", "na", "none"}:
        return None
    try:
        number = float(normalized)
    except ValueError:
        return None
    return int(number) if number.is_integer() else None


def _rows(path: str | Path, sheet_name: str, header_row: int) -> list[tuple[int, dict[str, Any]]]:
    from openpyxl import load_workbook

    workbook_path = Path(path)
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Workbook not found: {workbook_path.name}")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Missing sheet {sheet_name!r} in {workbook_path.name}")
        sheet = workbook[sheet_name]
        header_values = next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
        headers = [_clean(item) for item in header_values]
        output = []
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            record = {
                header: values[index]
                for index, header in enumerate(headers)
                if header and index < len(values)
            }
            if any(_clean(value) for value in record.values()):
                output.append((row_number, record))
        return output
    finally:
        workbook.close()


def build_import_payloads(
    *, category_index: str | Path, services: str | Path,
    hotels: str | Path, restaurants: str | Path,
) -> dict[str, list[dict[str, Any]]]:
    """Normalize workbook records without writing to Supabase."""
    category_path = Path(category_index)
    category_rows = _rows(category_path, "CATEGORIES", 1)
    mirror_rows = _rows(category_path, "BASEROW MIRROR", 1)
    categories_by_ref = {}
    enrich_by_ref = {}
    for source_row, row in category_rows:
        ref = _clean(row.get("Ref"))
        if not ref or ref.casefold() == "example":
            continue
        categories_by_ref[ref] = (source_row, row)
    for _, row in mirror_rows:
        ref = _clean(row.get("Ref"))
        if ref:
            enrich_by_ref[ref] = row

    category_payloads: list[dict[str, Any]] = []
    market_payloads: dict[str, dict[str, Any]] = {}
    market_category_payloads: dict[str, dict[str, Any]] = {}
    market_keys_by_name: dict[str, set[str]] = {}
    for ref, (source_row, row) in categories_by_ref.items():
        mirror = enrich_by_ref.get(ref, {})
        category_payloads.append({
            "source_ref": ref,
            "category_number": _integer(row.get("No") or mirror.get("No")),
            "category": _clean(row.get("Category") or mirror.get("Category")) or "Uncategorized",
            "group_name": _clean(row.get("Group") or mirror.get("Group")),
            "product_type": _clean(row.get("Product type") or mirror.get("Product type")) or "Unspecified",
            "priority": _clean(row.get("Priority")),
            "source_workbook": category_path.name,
            "source_sheet": "CATEGORIES",
            "source_row": source_row,
        })
        place_name = _clean(row.get("Market / place name") or mirror.get("Place"))
        city = _clean(row.get("City"))
        district = _clean(row.get("District or area"))
        address_value = _clean(row.get("Address if you know it"))
        address_en = _clean(mirror.get("Address EN"))
        address_cn = _clean(mirror.get("Address CN"))
        if address_value:
            if re.search(r"[\u4e00-\u9fff]", address_value):
                address_cn = address_cn or address_value
            else:
                address_en = address_en or address_value
        buyer_notes = _clean(row.get("What a buyer should know") or mirror.get("What you should know"))
        if place_name or address_en or address_cn:
            market_key = _key("market", place_name, city, district, address_en, address_cn)
            market_payloads.setdefault(market_key, {
                "source_key": market_key,
                "place_name": place_name,
                "city": city,
                "district": district,
                "address_en": address_en,
                "address_cn": address_cn,
                "buyer_notes": buyer_notes,
                "verification_status": _verification(row.get("Sure?")),
                "source_workbook": category_path.name,
                "source_sheet": "CATEGORIES",
                "source_row": source_row,
            })
            if place_name:
                market_keys_by_name.setdefault(place_name.casefold(), set()).add(market_key)
            link_key = _key("market-category", market_key, ref)
            market_category_payloads[link_key] = {
                "source_key": link_key,
                "category_ref": ref,
                "market_source_key": market_key,
                "source_workbook": category_path.name,
                "source_sheet": "CATEGORIES",
                "source_row": source_row,
            }

    service_path = Path(services)
    service_payloads = []
    for source_row, row in _rows(service_path, "Original Data", 1):
        service = _clean(row.get("Service"))
        if not service:
            continue
        source_key = _key("service", row.get("Company"), row.get("Phone / WeChat"), row.get("City"), service)
        service_payloads.append({
            "source_key": source_key,
            "service": service,
            "company": _clean(row.get("Company")),
            "company_cn": _clean(row.get("Company (CN)")),
            "contact_name": _clean(row.get("Person")),
            "phone_wechat": _clean(row.get("Phone / WeChat")),
            "other_numbers": _clean(row.get("Other numbers")),
            "email_qq": _clean(row.get("Email / QQ")),
            "city": _clean(row.get("City")),
            "address_en": _clean(row.get("Address (EN)")),
            "address_cn": _clean(row.get("Address (CN)")),
            "booth": _clean(row.get("Booth")),
            "ships_to": _clean(row.get("Ships to")),
            "hours": _clean(row.get("Hours")),
            "verification_status": _verification(row.get("Verified"), row.get("Source")),
            "notes": _clean(row.get("Notes")),
            "source": _clean(row.get("Source")),
            "source_workbook": service_path.name,
            "source_sheet": "Original Data",
            "source_row": source_row,
        })

    hotel_path = Path(hotels)
    hotel_payloads = []
    for source_row, row in _rows(hotel_path, "Original Data", 1):
        name = _clean(row.get("Hotel name"))
        if not name:
            continue
        source_key = _key("hotel", row.get("Area"), name)
        hotel_payloads.append({
            "source_key": source_key,
            "area": _clean(row.get("Area")),
            "hotel_name": name,
            "name_cn": _clean(row.get("Name (CN)")),
            "address_or_nearest_metro": _clean(row.get("Address / nearest metro")),
            "takes_foreigners": _boolean(row.get("Takes foreigners?")),
            "rough_price_per_night": _clean(row.get("Rough price a night")),
            "tier": _clean(row.get("Tier")),
            "notes": _clean(row.get("Notes")),
            "phone": _clean(row.get("Phone")),
            "verification_status": _verification(None, row.get("Source")),
            "source": _clean(row.get("Source")),
            "source_workbook": hotel_path.name,
            "source_sheet": "Original Data",
            "source_row": source_row,
        })

    restaurant_path = Path(restaurants)
    restaurant_payloads = []
    for source_row, row in _rows(restaurant_path, "Original Data", 1):
        name = _clean(row.get("Restaurant"))
        if not name:
            continue
        source_key = _key("restaurant", row.get("Area"), name)
        restaurant_payloads.append({
            "source_key": source_key,
            "area": _clean(row.get("Area")),
            "restaurant": name,
            "name_cn": _clean(row.get("Name (CN)")),
            "where_it_is": _clean(row.get("Where it is")),
            "food_type": _clean(row.get("Food type")),
            "halal": _boolean(row.get("Halal?")),
            "opening_hours": _clean(row.get("Opening hours")),
            "notes": _clean(row.get("Notes")),
            "phone": _clean(row.get("Phone")),
            "verification_status": _verification(None, row.get("Source")),
            "source": _clean(row.get("Source")),
            "source_workbook": restaurant_path.name,
            "source_sheet": "Original Data",
            "source_row": source_row,
        })

    city_guides = []
    for source_row, row in _rows(category_path, "CITY MAP", 4):
        city = _clean(row.get("City"))
        if not city:
            continue
        city_guides.append({
            "source_key": _key("city", city, row.get("Province")),
            "city": city,
            "province": _clean(row.get("Province")),
            "known_for": _clean(row.get("Known for")),
            "research_notes": _clean(row.get("What we are missing")),
            "verification_status": _verification(row.get("Right / wrong?")),
            "source_workbook": category_path.name,
            "source_sheet": "CITY MAP",
            "source_row": source_row,
        })

    category_mappings = []
    for source_row, row in _rows(category_path, "MAPPING", 4):
        old_category = _clean(row.get("Current category in the directory"))
        if not old_category:
            continue
        category_mappings.append({
            "source_key": _key("mapping", old_category, row.get("Goes to section"), source_row),
            "old_category": old_category,
            "entry_count": _integer(row.get("Entries")),
            "target_category_number": _integer(row.get("Goes to section")),
            "target_category_name": _clean(row.get("Section name")),
            "notes": _clean(row.get("Note")),
            "mapping_status": "needs_review" if _clean(row.get("Note")) else "suggested",
            "source_workbook": category_path.name,
            "source_sheet": "MAPPING",
            "source_row": source_row,
        })

    research_queue = []
    for sheet_name, issue_type in (("ADDRESS QUESTIONS", "address_conflict"), ("NAME THESE", "unnamed_market")):
        for source_row, row in _rows(category_path, sheet_name, 4):
            if not any(_clean(value) for value in row.values()):
                continue
            answer_field = "Your answer" if issue_type == "address_conflict" else "Market name"
            research_queue.append({
                "source_key": _key("research", category_path.name, sheet_name, source_row),
                "issue_type": issue_type,
                "source_payload": {key: _clean(value) for key, value in row.items() if _clean(value)},
                "status": "resolved" if _clean(row.get(answer_field)) else "open",
                "source_workbook": category_path.name,
                "source_sheet": sheet_name,
                "source_row": source_row,
            })

    contacts_payloads = []
    for source_row, row in _rows(category_path, "CONTACTS", 4):
        market_name = _clean(row.get("Market"))
        contact_fields = {
            "contact_name": _clean(row.get("Contact name")),
            "phone_wechat": _clean(row.get("Phone / WeChat")),
            "floor_or_stall": _clean(row.get("Floor or stall")),
            "best_time": _clean(row.get("Best time to find them")),
        }
        if not any(contact_fields.values()):
            continue
        matching_market_keys = market_keys_by_name.get((market_name or "").casefold(), set())
        if len(matching_market_keys) != 1:
            research_queue.append({
                "source_key": _key("research-contact", category_path.name, source_row),
                "issue_type": "contact_market_link_needs_review",
                "source_payload": {key: _clean(value) for key, value in row.items() if _clean(value)},
                "status": "open",
                "source_workbook": category_path.name,
                "source_sheet": "CONTACTS",
                "source_row": source_row,
            })
            continue
        market_key = next(iter(matching_market_keys))
        contacts_payloads.append({
            "source_key": _key("contact", market_key, *contact_fields.values()),
            "market_source_key": market_key,
            **contact_fields,
            "verification_status": "unknown",
            "source_workbook": category_path.name,
            "source_sheet": "CONTACTS",
            "source_row": source_row,
        })

    return {
        "directory_categories": category_payloads,
        "directory_markets": list(market_payloads.values()),
        "directory_market_categories": list(market_category_payloads.values()),
        "directory_contacts": contacts_payloads,
        "directory_service_providers": service_payloads,
        "directory_hotels": hotel_payloads,
        "directory_restaurants": restaurant_payloads,
        "directory_city_guides": city_guides,
        "directory_category_mappings": category_mappings,
        "directory_research_queue": research_queue,
    }


def import_payloads(payloads: dict[str, list[dict[str, Any]]]) -> dict[str, int]:
    """Write normalized batches. Caller must require an explicit apply flag."""
    from app.database.directory.categories import (
        upsert_categories, upsert_market_categories, upsert_markets,
    )
    from app.database.directory.city_guides import upsert_city_guides
    from app.database.directory.common import upsert_rows
    from app.database.directory.hotels import upsert_hotels
    from app.database.directory.restaurants import upsert_restaurants
    from app.database.directory.services import upsert_service_providers

    markets = payloads["directory_markets"]
    market_write_count = upsert_markets(markets)
    # Resolve source stable keys to database-generated UUIDs before inserting links.
    from app.database.client import supabase
    key_to_id = {}
    for offset in range(0, len(markets), 100):
        page = markets[offset:offset + 100]
        if not page:
            continue
        result = (
            supabase.table("directory_markets")
            .select("id,source_key")
            .in_("source_key", [row["source_key"] for row in page])
            .execute()
        )
        key_to_id.update({row["source_key"]: row["id"] for row in result.data or []})
    links = []
    for row in payloads["directory_market_categories"]:
        market_key = row["market_source_key"]
        market_id = key_to_id.get(market_key)
        if market_id:
            links.append({
                **{key: value for key, value in row.items() if key != "market_source_key"},
                "market_id": market_id,
            })
    contact_rows = []
    for row in payloads["directory_contacts"]:
        market_id = key_to_id.get(row["market_source_key"])
        if market_id:
            contact_rows.append({
                **{key: value for key, value in row.items() if key != "market_source_key"},
                "market_id": market_id,
            })
    counts = {
        "directory_categories": upsert_categories(payloads["directory_categories"]),
        "directory_markets": market_write_count,
        "directory_market_categories": upsert_market_categories(links),
        "directory_contacts": upsert_rows("directory_contacts", contact_rows, "source_key"),
        "directory_service_providers": upsert_service_providers(payloads["directory_service_providers"]),
        "directory_hotels": upsert_hotels(payloads["directory_hotels"]),
        "directory_restaurants": upsert_restaurants(payloads["directory_restaurants"]),
        "directory_city_guides": upsert_city_guides(payloads["directory_city_guides"]),
        "directory_category_mappings": upsert_rows("directory_category_mappings", payloads["directory_category_mappings"], "source_key"),
        "directory_research_queue": upsert_rows("directory_research_queue", payloads["directory_research_queue"], "source_key"),
    }
    return counts
