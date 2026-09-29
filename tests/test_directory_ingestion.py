from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from app.directory.ingestion import _integer, _verification, build_import_payloads


def make_workbook(path: Path, sheets: dict[str, list[list[object]]]) -> None:
    workbook = Workbook()
    first = True
    for name, rows in sheets.items():
        if first:
            sheet = workbook.active
            sheet.title = name
            first = False
        else:
            sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    workbook.save(path)
    workbook.close()


class DirectoryIngestionTests(unittest.TestCase):
    def test_not_verified_phrase_is_never_classified_as_verified(self) -> None:
        self.assertEqual(_verification("NOT verified"), "unverified")
        self.assertEqual(_verification(None, "Desk research, NOT verified"), "unverified")

    def test_numeric_placeholders_become_null_instead_of_database_errors(self) -> None:
        self.assertIsNone(_integer("-"))
        self.assertEqual(_integer(14), 14)
        self.assertEqual(_integer("1,200"), 1200)

    def test_workbooks_normalize_once_preserve_quality_and_provenance(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            category = root / "category.xlsx"
            service = root / "service.xlsx"
            hotel = root / "hotel.xlsx"
            restaurant = root / "restaurant.xlsx"
            make_workbook(category, {
                "CATEGORIES": [
                    ["Ref", "No", "Category", "Group", "Product type", "Priority", "Market / place name", "City", "District or area", "Address if you know it", "What a buyer should know", "Sure?"],
                    ["EXAMPLE", 1, "Example", "", "Example product", "P1", "", "", "", "", "", ""],
                    ["01-001", 1, "Fashion", "Bags", "Tote bags", "P1", "Sample Market", "Guangzhou", "Yuexiu", "中国广州市示例路1号", "Internal buyer note", ""],
                ],
                "BASEROW MIRROR": [
                    ["Ref", "No", "Category", "Group", "Product type", "Place", "Address EN", "Address CN", "What you should know"],
                    ["01-001", 1, "Fashion", "Bags", "Tote bags", "Sample Market", "", "中国广州市示例路1号", "Internal buyer note"],
                ],
                "CITY MAP": [
                    ["title"], ["instructions"], [],
                    ["City", "Province", "Known for", "Right / wrong?", "What we are missing"],
                    ["Guangzhou", "Guangdong", "Fashion", "", "Needs review"],
                ],
                "MAPPING": [
                    ["title"], ["instructions"], [],
                    ["Current category in the directory", "Entries", "Goes to section", "Section name", "Note"],
                    ["Old Bags", 2, 1, "Fashion", "Check mapping"],
                ],
                "ADDRESS QUESTIONS": [
                    ["title"], ["instructions"], [],
                    ["Market", "The question", "Your answer", "How sure?"],
                    ["Sample Market", "Which address?", "", ""],
                ],
                "NAME THESE": [
                    ["title"], ["instructions"], [],
                    ["Address", "District", "Sells", "Market name", "Anything else worth knowing"],
                    ["Road 1", "Yuexiu", "Bags", "", ""],
                ],
                "CONTACTS": [
                    ["title"], ["instructions"], [],
                    ["Market", "Contact name", "Phone / WeChat", "Floor or stall", "Best time to find them"],
                    ["Sample Market", "Contact Person", "+86 100", "A1", "Morning"],
                ],
            })
            make_workbook(service, {"Original Data": [
                ["Service", "Company", "Company (CN)", "Person", "Phone / WeChat", "Other numbers", "Email / QQ", "City", "Address (EN)", "Address (CN)", "Booth", "Ships to", "Hours", "Verified", "Notes", "Source"],
                ["Freight forwarder", "Company A", "公司A", "Person A", "+86 100", None, None, "Guangzhou", "Street", "街道", "A1", "Senegal", None, "Y", "Private", "Verified card"],
            ]})
            make_workbook(hotel, {"Original Data": [
                ["Area", "Hotel name", "Name (CN)", "Address / nearest metro", "Takes foreigners?", "Rough price a night", "Tier", "Notes", "Phone", "Source"],
                ["Pazhou", "Hotel A", None, "Metro A", None, None, "Budget", "Internal", "+86 1", "Desk research, NOT verified"],
            ]})
            make_workbook(restaurant, {"Original Data": [
                ["Area", "Restaurant", "Name (CN)", "Where it is", "Food type", "Halal?", "Opening hours", "Notes", "Phone", "Source"],
                ["Xiaobei", "Restaurant A", None, "Street A", "West African", "No", None, "Internal", "+86 2", "Desk research, NOT verified"],
            ]})

            args = dict(category_index=category, services=service, hotels=hotel, restaurants=restaurant)
            payloads = build_import_payloads(**args)
            repeated = build_import_payloads(**args)

        self.assertEqual(len(payloads["directory_categories"]), 1)
        self.assertEqual(payloads["directory_categories"][0]["source_ref"], "01-001")
        self.assertEqual(len(payloads["directory_markets"]), 1)
        self.assertEqual(payloads["directory_markets"][0]["verification_status"], "unknown")
        self.assertEqual(payloads["directory_markets"][0]["address_cn"], "中国广州市示例路1号")
        self.assertEqual(len(payloads["directory_contacts"]), 1)
        self.assertEqual(payloads["directory_contacts"][0]["verification_status"], "unknown")
        self.assertEqual(payloads["directory_service_providers"][0]["verification_status"], "verified")
        self.assertEqual(payloads["directory_hotels"][0]["verification_status"], "unverified")
        self.assertIsNone(payloads["directory_hotels"][0]["takes_foreigners"])
        self.assertFalse(payloads["directory_restaurants"][0]["halal"])
        self.assertEqual(payloads["directory_restaurants"][0]["verification_status"], "unverified")
        self.assertEqual(len(payloads["directory_research_queue"]), 2)
        self.assertEqual(
            [r["source_key"] for r in payloads["directory_markets"]],
            [r["source_key"] for r in repeated["directory_markets"]],
        )


if __name__ == "__main__":
    unittest.main()
